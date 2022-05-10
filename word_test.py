from docx import Document
import pandas as df
import protocol_fish_counting_SQL as fc
import initial_variables as init_var

# i am not sure how you are getting your data, but you said it is a
# pandas data frame
df_age_struct = fc.get_age_struct_from_sql(
            sql_server=init_var.current_sql_server,
            bio_space=init_var.current_bio_space,
            cur_year=init_var.current_cur_year,
            water_body='el_watrer_body',
            sql_server_name=init_var.current_sql_server_name,
            data_base_name=init_var.current_data_base_name
        )
# open an existing document
doc = Document('./test.docx')
# add a table to the end and create a reference variable
# extra row is so we can add the header row
t = doc.add_table(df.shape[0]+1, df.shape[1])
# add the header rows.
for j in range(df.shape[-1]):
    t.cell(0, j).text = df.columns[j]
# add the rest of the data frame
for i in range(df.shape[0]):
    for j in range(df.shape[-1]):
        t.cell(i+1, j).text = str(df.values[i,j])
# save the doc
doc.save('./test.docx')
